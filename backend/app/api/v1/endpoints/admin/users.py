from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from uuid import UUID

from fastapi import APIRouter, File, HTTPException, Query, UploadFile, status
from sqlalchemy import func, select
from sqlalchemy.orm import selectinload

from app.core.dependencies import AdminUser, DbSession
from app.core.security import hash_password
from app.models.assignment import StudentSectionAssignment, TeacherSectionAssignment
from app.models.class_ import SchoolClass
from app.models.section import Section
from app.models.user import User, UserRole
from app.repositories.user_repo import UserRepository
from app.schemas.user import UpdateRoleIn, UpdateStatusIn, UserOut

router = APIRouter()


def _user_base(u: User) -> dict:
    return {
        "id": str(u.id),
        "full_name": u.full_name,
        "email": u.email,
        "role": u.role.value,
        "is_active": u.is_active,
        "created_at": u.created_at,
        "assigned_section": None,
        "assigned_sections": None,
    }


@router.get("")
async def list_users(
    db: DbSession, _: AdminUser,
    page: int = Query(1, ge=1), limit: int = Query(20, ge=1, le=100),
    role: UserRole | None = None, is_active: bool | None = None,
    search: str | None = None,
):
    repo = UserRepository(db)
    users, total = await repo.get_paginated(page=page, limit=limit, role=role, is_active=is_active, search=search)

    if not users:
        return {"data": [], "pagination": {"total": total, "page": page, "limit": limit}}

    student_ids = [u.id for u in users if u.role == UserRole.student]
    teacher_ids = [u.id for u in users if u.role == UserRole.teacher]

    student_map: dict = {}
    if student_ids:
        rows = (await db.execute(
            select(StudentSectionAssignment)
            .where(
                StudentSectionAssignment.student_id.in_(student_ids),
                StudentSectionAssignment.is_active == True,
            )
            .options(selectinload(StudentSectionAssignment.section).selectinload(Section.school_class))
        )).scalars().all()
        for a in rows:
            sec = a.section
            cls = sec.school_class if sec else None
            student_map[a.student_id] = {
                "section_id": str(a.section_id),
                "section_name": sec.name if sec else "",
                "class_name": cls.name if cls else "",
            }

    teacher_map: dict = {}
    if teacher_ids:
        rows = (await db.execute(
            select(TeacherSectionAssignment)
            .where(
                TeacherSectionAssignment.teacher_id.in_(teacher_ids),
                TeacherSectionAssignment.is_active == True,
            )
            .options(selectinload(TeacherSectionAssignment.section).selectinload(Section.school_class))
        )).scalars().all()
        for a in rows:
            sec = a.section
            cls = sec.school_class if sec else None
            teacher_map.setdefault(a.teacher_id, []).append({
                "section_id": str(a.section_id),
                "section_name": sec.name if sec else "",
                "class_name": cls.name if cls else "",
            })

    result = []
    for u in users:
        d = _user_base(u)
        if u.role == UserRole.student:
            d["assigned_section"] = student_map.get(u.id)
        elif u.role == UserRole.teacher:
            d["assigned_sections"] = teacher_map.get(u.id, [])
        result.append(d)

    return {"data": result, "pagination": {"total": total, "page": page, "limit": limit}}


def _parse_csv_rows(contents: bytes) -> list[dict]:
    text = contents.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [{k.strip().lower(): (v or "").strip() for k, v in row.items() if k} for row in reader]


def _parse_excel_rows(contents: bytes) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(contents), data_only=True)
    ws = wb.active
    headers: list[str] = []
    rows: list[dict] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c).strip().lower() if c is not None else "" for c in row]
        else:
            rows.append({
                headers[j]: str(row[j]).strip() if row[j] is not None else ""
                for j in range(len(headers))
            })
    return rows


@router.post("/import", status_code=status.HTTP_201_CREATED)
async def import_users(
    db: DbSession,
    admin: AdminUser,
    file: UploadFile = File(...),
    role: UserRole = Query(..., description="Role to assign: student or teacher"),
):
    if role == UserRole.admin:
        raise HTTPException(status_code=400, detail="Cannot import admin users")

    contents = await file.read()
    fname = (file.filename or "").lower()

    if fname.endswith(".xlsx") or fname.endswith(".xls"):
        rows = _parse_excel_rows(contents)
    else:
        rows = _parse_csv_rows(contents)

    if not rows:
        raise HTTPException(status_code=400, detail="File is empty or has no data rows")

    created_users: list[dict] = []
    errors: list[dict] = []

    for idx, row in enumerate(rows, start=2):
        name = row.get("name", "").strip()
        username = (row.get("email") or row.get("username") or "").strip()
        password = row.get("password", "").strip()
        class_name = row.get("class", "").strip()
        section_name = row.get("section", "").strip()

        if not all([name, username, password, class_name, section_name]):
            errors.append({"row": idx, "error": "Missing required fields (Name, Email/Username, Password, Class, Section)"})
            continue

        email = (username if "@" in username else f"{username}@school.local").lower()

        cls = (await db.execute(
            select(SchoolClass).where(func.lower(SchoolClass.name) == class_name.lower())
        )).scalar_one_or_none()
        if not cls:
            errors.append({"row": idx, "error": f"Class '{class_name}' not found"})
            continue

        section = (await db.execute(
            select(Section).where(
                Section.class_id == cls.id,
                func.lower(Section.name) == section_name.lower(),
                Section.is_active == True,
            )
        )).scalar_one_or_none()
        if not section:
            errors.append({"row": idx, "error": f"Section '{section_name}' not found in class '{class_name}'"})
            continue

        existing = (await db.execute(select(User).where(User.email == email))).scalar_one_or_none()
        if existing:
            errors.append({"row": idx, "error": f"Email '{email}' already exists"})
            continue

        user = User(
            full_name=name,
            email=email,
            password_hash=hash_password(password),
            role=role,
            is_active=True,
        )
        db.add(user)
        await db.flush()

        now = datetime.now(timezone.utc)
        if role == UserRole.student:
            db.add(StudentSectionAssignment(
                student_id=user.id,
                section_id=section.id,
                assigned_by=admin.id,
                is_active=True,
                assigned_at=now,
            ))
        else:
            db.add(TeacherSectionAssignment(
                teacher_id=user.id,
                section_id=section.id,
                assigned_by=admin.id,
                is_active=True,
                assigned_at=now,
            ))

        await db.flush()
        created_users.append({"name": name, "email": email})

    return {"created": len(created_users), "errors": errors, "users": created_users}


@router.get("/{userId}")
async def get_user(userId: UUID, db: DbSession, _: AdminUser):
    repo = UserRepository(db)
    user = await repo.get_by_id(userId)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return UserOut.model_validate(user)


@router.patch("/{userId}/role")
async def update_role(userId: UUID, body: UpdateRoleIn, db: DbSession, admin: AdminUser):
    if userId == admin.id:
        raise HTTPException(status_code=400, detail="Cannot change your own role")
    repo = UserRepository(db)
    user = await repo.get_by_id(userId)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    if user.role == UserRole.admin:
        raise HTTPException(status_code=400, detail="Cannot change admin role")
    user.role = body.role
    user.is_active = True
    await repo.save(user)
    return UserOut.model_validate(user)


@router.patch("/{userId}/status")
async def update_status(userId: UUID, body: UpdateStatusIn, db: DbSession, admin: AdminUser):
    if userId == admin.id:
        raise HTTPException(status_code=400, detail="Cannot deactivate your own account")
    repo = UserRepository(db)
    user = await repo.get_by_id(userId)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    user.is_active = body.is_active
    await repo.save(user)
    return UserOut.model_validate(user)
